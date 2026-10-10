import io
from typing import List, Dict, Any, Optional
from sqlalchemy.orm import Session
import pypdf
import docx

from app.models.models import Document, DocumentChunk
from app.services.embedding_service import embedding_service
from app.services.llm_service import llm_service

class RAGService:
    @staticmethod
    def chunk_text(text: str, chunk_size: int = 500, chunk_overlap: int = 80) -> List[str]:
        words = text.split()
        if not words:
            return []
        
        chunks = []
        start = 0
        while start < len(words):
            end = start + chunk_size
            chunk_str = " ".join(words[start:end])
            if chunk_str.strip():
                chunks.append(chunk_str.strip())
            start += (chunk_size - chunk_overlap)
        return chunks

    @classmethod
    async def index_document(
        cls,
        db: Session,
        user_id: str,
        filename: str,
        file_bytes: bytes,
        topic: Optional[str] = None
    ) -> Document:
        # Extract text based on file type
        filename_lower = filename.lower()
        text = ""
        file_type = "txt"
        
        if filename_lower.endswith(".pdf"):
            file_type = "pdf"
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        elif filename_lower.endswith(".docx"):
            file_type = "docx"
            doc = docx.Document(io.BytesIO(file_bytes))
            for p in doc.paragraphs:
                if p.text:
                    text += p.text + "\n"
        else:
            file_type = "text"
            try:
                text = file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                text = file_bytes.decode("latin-1", errors="ignore")

        # Create Document record
        doc_record = Document(
            user_id=user_id,
            filename=filename,
            topic=topic or "General Preparation",
            file_type=file_type,
            char_count=len(text)
        )
        db.add(doc_record)
        db.flush()

        # Chunk text
        chunks = cls.chunk_text(text, chunk_size=300, chunk_overlap=50)
        for idx, chunk_text in enumerate(chunks):
            # Compute real embedding for chunk
            embedding = await embedding_service.get_embedding(chunk_text)
            
            chunk_rec = DocumentChunk(
                document_id=doc_record.id,
                user_id=user_id,  # Strict user isolation
                chunk_index=idx,
                content=chunk_text,
                embedding=embedding,
                token_count=len(chunk_text.split())
            )
            db.add(chunk_rec)

        db.commit()
        db.refresh(doc_record)
        return doc_record

    @classmethod
    async def query_user_documents(
        cls,
        db: Session,
        user_id: str,
        query: str,
        limit: int = 4
    ) -> List[Dict[str, Any]]:
        # Enforce strict user isolation: only fetch chunks belonging to this user
        user_chunks = (
            db.query(DocumentChunk, Document.filename)
            .join(Document, DocumentChunk.document_id == Document.id)
            .filter(DocumentChunk.user_id == user_id)
            .all()
        )

        if not user_chunks:
            return []

        query_embedding = await embedding_service.get_embedding(query)
        scored_results = []

        for chunk, filename in user_chunks:
            if chunk.embedding:
                score = embedding_service.cosine_similarity(query_embedding, chunk.embedding)
                scored_results.append({
                    "chunk_index": chunk.chunk_index,
                    "content": chunk.content,
                    "similarity_score": round(score, 4),
                    "document_filename": filename
                })

        # Sort descending by similarity score
        scored_results.sort(key=lambda x: x["similarity_score"], reverse=True)
        return scored_results[:limit]

rag_service = RAGService()
